"""Owned hosted fixture only: independent producer, never a native IPC mock."""
import hashlib, json, os, sys, zipfile
from pathlib import Path
import openpyxl
from openpyxl.worksheet.filters import AutoFilter

for key, expected in {'GITHUB_ACTIONS':'true','RUNNER_ENVIRONMENT':'github-hosted','RUNNER_OS':'Windows','LONGEDIT_R5I_DISPOSABLE':'1'}.items():
    assert os.environ.get(key) == expected, 'Disposable hosted XLSX fixture only'
mode, owner_arg, file_arg, output_arg = sys.argv[1:]
owner, file, output = Path(owner_arg).resolve(), Path(file_arg).resolve(), Path(output_arg).resolve()
assert owner.is_relative_to(Path('C:/LongEditR5ILibrary').resolve()) and (owner / '.v119-installed-owner').is_file()
assert file.is_relative_to(owner) and file.suffix == '.xlsx' and not output.exists()
assert openpyxl.__version__ == '3.1.5'
sha=lambda file: hashlib.sha256(file.read_bytes()).hexdigest()
if mode == 'create':
    assert not file.exists()
    workbook=openpyxl.Workbook(); sheet=workbook.active; sheet.title='任务'
    for row in [['任务','状态','编号','数值'],['会议\r\n旧备注','可见','00123',0.1],['不要修改','隐藏','9007199254740993124',0.2],['核对文档','可见','9007199254740993123',3],['=保留文本','可见','00007','=D2+D3']]: sheet.append(row)
    sheet['A5'].data_type='s';sheet.auto_filter=AutoFilter(ref='A1:D5');sheet['A2'].alignment=openpyxl.styles.Alignment(wrap_text=True)
    sheet.column_dimensions['A'].width=28; sheet.column_dimensions['C'].width=30
    workbook.create_sheet('不要修改')['A1']='其他工作表保持原样'
    workbook.save(file);workbook.close()
    output.write_bytes(file.read_bytes()) # Immutable before-image, owned fixture.
elif mode == 'inspect':
    assert file.stat().st_size < 32*1024*1024
    workbook=openpyxl.load_workbook(file, data_only=False, keep_links=False);sheet=workbook['任务']
    cells={cell.coordinate:str(cell.value) if cell.value is not None else '' for row in sheet for cell in row}
    types={cell.coordinate:cell.data_type for row in sheet for cell in row}
    expected={'C2':'00124','C3':'9007199254740993124','C4':'9007199254740993123','A5':'=保留文本','A6':'=保留文本','D4':'0.3','D6':'=D3+D4'}
    for address,value in expected.items(): assert cells[address] == value, (address,cells[address],value)
    for address in ['C2','C3','C4','A5','A6']: assert types[address]=='s', (address,types[address])
    assert types['D4']=='n' and types['D6']=='f' and sheet['A3'].value=='不要修改'
    assert cells['A2'].replace('\r\n','\n')=='会议\n备注'
    assert workbook['不要修改']['A1'].value=='其他工作表保持原样'
    before=owner/'原始 工作簿.xlsx';preserved=[]
    with zipfile.ZipFile(before) as original, zipfile.ZipFile(file) as saved:
        for name in ['xl/worksheets/sheet2.xml','docProps/core.xml','docProps/app.xml','xl/_rels/workbook.xml.rels','[Content_Types].xml']:
            assert original.read(name)==saved.read(name), name
            preserved.append({'part':name,'sha256':hashlib.sha256(saved.read(name)).hexdigest()})
    producer=owner/'独立生产者 另存.xlsx';assert not producer.exists();workbook.save(producer);workbook.close()
    reopened=openpyxl.load_workbook(producer,data_only=False,keep_links=False)
    for address,value in expected.items(): assert str(reopened['任务'][address].value)==value
    reopened.close()
    report={'status':'passed','nativeBridgeMocked':False,'sourceUserContentIncluded':False,'reader':'openpyxl-3.1.5-hash-pinned','sha256':sha(file),'beforeSha256':sha(before),'cells':cells,'types':types,'preservedParts':preserved,'producerReopen':{'status':'passed','sha256':sha(producer),'formulaCalculated':False,'officeVisualAcceptance':False}}
    output.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
else:
    raise AssertionError('Unknown fixture mode')
