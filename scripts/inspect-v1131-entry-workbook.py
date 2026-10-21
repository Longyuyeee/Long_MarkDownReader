import sys,json,hashlib,zipfile
from pathlib import Path
from openpyxl import load_workbook
before,after=map(Path,sys.argv[1:3])
a,b=load_workbook(before),load_workbook(after)
assert a.sheetnames==b.sheetnames==['录入','保留','Options']
s=b['录入'];assert s['A2'].value==1.5 and s['B2'].value==.1 and s['C2'].value=='😀字' and str(s['D2'].value).startswith('2026-10-23') and s['A5'].value=='未保存保护三一'
assert s['E4'].value=='=A4' and s.row_dimensions[3].hidden and s['A2'].font.bold
for name in ['保留','Options']:
 assert list(a[name].values)==list(b[name].values)
assert [r.__dict__ for r in a['录入'].data_validations.dataValidation]==[r.__dict__ for r in s.data_validations.dataValidation]
with zipfile.ZipFile(before) as za,zipfile.ZipFile(after) as zb:
 assert sorted(za.namelist())==sorted(zb.namelist())
 changed=[n for n in za.namelist() if za.read(n)!=zb.read(n)]
 assert 'xl/worksheets/sheet1.xml' in changed and set(changed)<=set(['xl/worksheets/sheet1.xml','xl/workbook.xml','xl/styles.xml'])
for row in [3,4]:
 for col in range(1,6):assert a['录入'].cell(row,col).value==s.cell(row,col).value
print(json.dumps({'status':'passed','beforeSha256':hashlib.sha256(before.read_bytes()).hexdigest(),'sha256':hashlib.sha256(after.read_bytes()).hexdigest(),'changedParts':changed,'sheetNames':b.sheetnames,'values':{'A2':s['A2'].value,'B2':s['B2'].value,'C2':s['C2'].value,'D2':str(s['D2'].value),'A5':s['A5'].value,'E4':s['E4'].value}},ensure_ascii=True))
