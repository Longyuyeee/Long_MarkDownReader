import os,sys
from pathlib import Path
from datetime import date
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
assert os.environ.get('GITHUB_ACTIONS')=='true' and os.environ.get('LONGEDIT_R5I_DISPOSABLE')=='1' and os.environ.get('RUNNER_ENVIRONMENT')=='github-hosted'
owner=Path(sys.argv[1]).resolve();target=Path(sys.argv[2]).resolve()
owner.relative_to(Path('C:/LongEditR5ILibrary').resolve());target.relative_to(owner)
assert target.suffix=='.xlsx' and not target.exists() and (owner/'.v119-installed-owner').exists()
book=Workbook();sheet=book.active;sheet.title='录入'
for row in [['数量','单价','备注','日期','区域列表'],[1.5,.1,'太长的备注',date(2026,10,21),'甲'],[99,99,'隐藏原文',date(2025,1,1),'乙'],[2,.2,'好',date(2026,10,22),'=A4'],['继续录入',None,None,None,None]]:sheet.append(row)
sheet.row_dimensions[3].hidden=True
for col in 'ABCDE':sheet.column_dimensions[col].width=12
sheet.auto_filter.ref='A1:E4'
for kind,source,end,col in [('whole','1','10','A'),('decimal','0.1','1','B'),('textLength','0','3','C'),('date','46316','46319','D'),('list','Options!$A$1:$A$3',None,'E')]:
 rule=DataValidation(type=kind,operator='between' if end else None,formula1=source,formula2=end,allow_blank=False)
 rule.showErrorMessage=True;sheet.add_data_validation(rule);rule.add(f'{col}2:{col}4')
sheet['A2'].font=sheet['A2'].font.copy(bold=True)
book.create_sheet('保留').append(['00123','9007199254740993123','=1+2'])
book.create_sheet('Options').append(['甲','乙','丙'])
book.save(target)
